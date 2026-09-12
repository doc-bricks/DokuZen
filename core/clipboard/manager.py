# -*- coding: utf-8 -*-
"""Clipboard-zu-Datei Manager für DokuZen.

Ermöglicht das Speichern von Zwischenablage-Inhalten in verschiedenen Formaten
(TXT, MD, PDF, DOCX, RTF) mit Struktur- und Überschriften-Erkennung sowie
atomarer Dateispeicherung.
"""

from __future__ import annotations

import textwrap
from datetime import datetime
from pathlib import Path
from typing import Optional

from core.clipboard.naming import slugify_filename, unique_path
from core.clipboard.structure import analyze_structure
from utils.logger import get_logger

_logger = get_logger(__name__)

# ── Optionale Abhängigkeiten (Lazy / Graceful Fallback) ───────────────────
_HAS_PYPERCLIP = False
_HAS_REPORTLAB = False
_HAS_DOCX = False

try:
    import pyperclip as _pyperclip
    _HAS_PYPERCLIP = True
except ImportError:
    _pyperclip = None

try:
    from reportlab.lib.pagesizes import A4 as _A4
    from reportlab.pdfbase import pdfmetrics as _pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as _TTFont
    from reportlab.pdfgen import canvas as _rl_canvas
    _HAS_REPORTLAB = True
except ImportError:
    _A4 = (595.27, 841.89)
    _rl_canvas = None
    _pdfmetrics = None
    _TTFont = None

try:
    from docx import Document as _Document
    _HAS_DOCX = True
except ImportError:
    _Document = None

# ── Layout-Konstanten ─────────────────────────────────────────────────────
FORMATS = ["TXT", "MD", "PDF", "DOCX", "RTF"]

_PDF_MARGIN = 40
_PDF_HEADING_FONTSIZE = 14
_PDF_BODY_FONTSIZE = 11
_PDF_LINE_SPACING = 4
_PDF_MIN_WRAP_WIDTH = 20


class ClipboardManager:
    """Verwaltet das Einlesen der Zwischenablage und Speichern als formatierte Datei.

    Unterstützt:
    - Textanalyse und Überschriftenerkennung (Design-Modus)
    - Automatische Namensgenerierung aus Inhalt / Überschrift
    - Eindeutige Zieldateipfade (Kollisionsvermeidung)
    - Formate: TXT, MD, PDF, DOCX, RTF
    - Atomare Dateierstellung (.tmp-Datei -> replace)
    """

    FORMATS = FORMATS

    def __init__(self, default_dir: Optional[str] = None) -> None:
        if default_dir:
            self.default_dir = Path(default_dir)
        else:
            self.default_dir = Path.home() / "Documents" / "DokuZen" / "Spawned"

    def get_clipboard_text(self) -> str:
        """Liest den aktuellen Rohtext aus der Zwischenablage."""
        if not _HAS_PYPERCLIP:
            raise RuntimeError(
                "pyperclip ist nicht installiert. Bitte mit 'pip install pyperclip' nachinstallieren."
            )
        try:
            return (_pyperclip.paste() or "").strip()
        except Exception as exc:
            raise RuntimeError(
                f"Zwischenablage konnte nicht gelesen werden: {exc}"
            ) from exc

    def save_from_clipboard(
        self,
        output_dir: Optional[str] = None,
        fmt: str = "TXT",
        design: bool = True,
        custom_text: Optional[str] = None,
    ) -> str:
        """Liest die Zwischenablage (oder optional custom_text) und speichert den Inhalt.

        Parameters
        ----------
        output_dir:
            Zielverzeichnis (Standard: self.default_dir).
        fmt:
            Ausgabeformat ('TXT', 'MD', 'PDF', 'DOCX', 'RTF').
        design:
            Wenn True, wird der Text strukturiert und Überschriften formatiert.
        custom_text:
            Optionaler Text zur Umgehung der System-Zwischenablage (z.B. für Tests).

        Returns
        -------
        str
            Absoluter Pfad der erstellten Datei.
        """
        raw = custom_text.strip() if custom_text is not None else self.get_clipboard_text()
        if not raw:
            raise RuntimeError("Zwischenablage oder Quelltext ist leer.")

        fmt_upper = fmt.upper()
        if fmt_upper not in FORMATS:
            raise ValueError(
                f"Nicht unterstütztes Format: '{fmt}'. Gültig sind: {', '.join(FORMATS)}"
            )

        target_dir = Path(output_dir) if output_dir else self.default_dir
        target_dir.mkdir(parents=True, exist_ok=True)

        if design:
            struct = analyze_structure(raw)
        else:
            struct = [{"type": "paragraph", "text": raw}]

        slug = slugify_filename(raw)
        ext = fmt_upper.lower()
        target_path = unique_path(target_dir / slug, ext)

        path_str = str(target_path)
        if fmt_upper == "TXT":
            self.save_txt(path_str, struct)
        elif fmt_upper == "MD":
            self.save_md(path_str, struct)
        elif fmt_upper == "PDF":
            self.save_pdf(path_str, struct)
        elif fmt_upper == "DOCX":
            self.save_docx(path_str, struct)
        elif fmt_upper == "RTF":
            self.save_rtf(path_str, struct)

        _logger.info(f"Clipboard erfolgreich gespeichert: {path_str} (Format: {fmt_upper})")
        return path_str

    # ── Format-Speicherfunktionen (atomar via .tmp -> replace) ─────────────

    @staticmethod
    def save_txt(path: str, struct: list[dict]) -> None:
        """Speichert strukturierten Text als Plain-Text-Datei."""
        lines: list[str] = []
        for el in struct:
            if el.get("type") == "heading":
                lines.append(f"# {el['text']}")
            else:
                lines.append(el["text"])
            lines.append("")

        content = "\n".join(lines)
        tmp_path = Path(path).with_suffix(".tmp")
        try:
            tmp_path.write_text(content, encoding="utf-8")
            tmp_path.replace(Path(path))
        except Exception:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise

    @staticmethod
    def save_md(path: str, struct: list[dict], title: Optional[str] = None) -> None:
        """Speichert strukturierten Text als Markdown-Dokument."""
        lines: list[str] = []
        doc_title = title
        if not doc_title and struct and struct[0].get("type") == "heading":
            doc_title = struct[0]["text"]

        lines.append(f"# {doc_title or 'Clipboard-Export'}")
        lines.append(f"> Erstellt am: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("")
        lines.append("---")
        lines.append("")

        for el in struct:
            if el.get("type") == "heading":
                level = el.get("level", 2)
                prefix = "#" * max(2, min(level, 6))
                lines.append(f"{prefix} {el['text']}")
            else:
                lines.append(el["text"])
            lines.append("")

        content = "\n".join(lines)
        tmp_path = Path(path).with_suffix(".tmp")
        try:
            tmp_path.write_text(content, encoding="utf-8")
            tmp_path.replace(Path(path))
        except Exception:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise

    @staticmethod
    def save_pdf(path: str, struct: list[dict]) -> None:
        """Speichert strukturierten Text als formatiertes PDF-Dokument."""
        if not _HAS_REPORTLAB:
            raise RuntimeError(
                "reportlab ist nicht installiert. PDF-Export nicht möglich."
            )

        w, h = _A4
        margin = _PDF_MARGIN

        font = "Helvetica"
        is_unicode_font = False
        try:
            if _pdfmetrics and _TTFont:
                _pdfmetrics.registerFont(_TTFont("ArialUnicode", "arialuni.ttf"))
                font = "ArialUnicode"
                is_unicode_font = True
        except Exception:
            font = "Helvetica"
            is_unicode_font = False

        tmp_path = Path(path).with_suffix(".tmp")
        try:
            c = _rl_canvas.Canvas(str(tmp_path), pagesize=_A4)
            y = h - margin

            for el in struct:
                is_head = el.get("type") == "heading"
                sz = _PDF_HEADING_FONTSIZE if is_head else _PDF_BODY_FONTSIZE
                c.setFont(font, sz)

                wrap_width = max(
                    _PDF_MIN_WRAP_WIDTH,
                    int((w - 2 * margin) / (sz * 0.5)),
                )

                for raw_line in textwrap.wrap(el["text"], wrap_width):
                    if y < margin:
                        c.showPage()
                        c.setFont(font, sz)
                        y = h - margin

                    line = raw_line if is_unicode_font else raw_line.encode("latin-1", "replace").decode("latin-1")
                    c.drawString(margin, y, line)
                    y -= sz + _PDF_LINE_SPACING

                y -= sz

            c.save()
            tmp_path.replace(Path(path))
        except Exception:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise

    @staticmethod
    def save_docx(path: str, struct: list[dict]) -> None:
        """Speichert strukturierten Text als Microsoft Word (DOCX) Dokument."""
        if not _HAS_DOCX:
            raise RuntimeError(
                "python-docx ist nicht installiert. DOCX-Export nicht möglich."
            )

        tmp_path = Path(path).with_suffix(".tmp")
        try:
            doc = _Document()
            for el in struct:
                if el.get("type") == "heading":
                    level = el.get("level", 1)
                    doc.add_heading(el["text"], level=min(level, 3))
                else:
                    doc.add_paragraph(el["text"])

            doc.save(str(tmp_path))
            tmp_path.replace(Path(path))
        except Exception:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise

    @staticmethod
    def _rtf_escape(text: str) -> str:
        """Escaped Zeichen und Zeilenumbrüche für RTF-Dokumente inklusive Unicode."""
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        parts: list[str] = []
        for ch in normalized:
            cp = ord(ch)
            if ch == "\n":
                parts.append(r"\line ")
            elif cp < 128:
                if ch == '\\':
                    parts.append(r'\\')
                elif ch == '{':
                    parts.append(r'\{')
                elif ch == '}':
                    parts.append(r'\}')
                else:
                    parts.append(ch)
            elif cp <= 255:
                parts.append(f"\\'{cp:02x}")
            elif cp <= 0xFFFF:
                signed = cp if cp < 32768 else cp - 65536
                parts.append(f"\\u{signed}?")
            else:
                cp_sub = cp - 0x10000
                high = 0xD800 + (cp_sub >> 10)
                low = 0xDC00 + (cp_sub & 0x3FF)
                high_signed = high if high < 32768 else high - 65536
                low_signed = low if low < 32768 else low - 65536
                parts.append(f"\\u{high_signed}?\\u{low_signed}?")
        return "".join(parts)

    @classmethod
    def save_rtf(cls, path: str, struct: list[dict]) -> None:
        """Speichert strukturierten Text als Rich Text Format (RTF) Dokument."""
        rtf_lines: list[str] = [r"{\rtf1\ansi\deff0"]
        for el in struct:
            escaped = cls._rtf_escape(el["text"])
            if el.get("type") == "heading":
                rtf_lines.append(r"\b " + escaped + r" \b0")
            else:
                rtf_lines.append(escaped)
            rtf_lines.append(r"\par")
        rtf_lines.append("}")

        content = "\n".join(rtf_lines).encode("ascii")
        tmp_path = Path(path).with_suffix(".tmp")
        try:
            tmp_path.write_bytes(content)
            tmp_path.replace(Path(path))
        except Exception:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise
